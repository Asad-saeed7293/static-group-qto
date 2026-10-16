import os
import sys
import base64
import hashlib
import hmac
import json
import multiprocessing

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from qto import store
from qto.boq import mappings
from qto.boq.compare import compare
from qto.review.queue import build_queue


SETTINGS = {"tolerances": [{"max_abs_pct": 2.0, "status": "OK"},
                           {"max_abs_pct": 5.0, "status": "Check"}],
            "tolerance_above_status": "Review Required", "verified_requires_cross_checks": True}


def _row(source_row=8, **overrides):
    row = {"row": source_row, "source_row": source_row, "source_sheet": "BOQ", "item_no": "A-01",
           "description": "Reinforced concrete footing F3", "element": "", "floor": "",
           "unit": "m3", "quantity": 3.6, "section": "Footings", "issues": [],
           "source_item_no": "A-01", "source_description": "Reinforced concrete footing F3",
           "source_unit": "m3", "source_quantity": 3.6}
    row.update(overrides)
    return row


def _line(key="line-footing", item_no="02.01", description="Reinforced concrete footing F3",
          unit="m3", category="Footings", quantity=3.6):
    return {"match_key": key, "item_no": item_no, "mark_norm": "", "element_type": "footing",
            "floor": "FDN", "description": description, "quantity": quantity, "unit": unit,
            "status": "OK", "category": category, "element_uids": ["F3-1"], "qids": ["F3:concrete"],
            "calculation_evidence": []}


@pytest.fixture
def mapping_project(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "DATA", str(tmp_path))
    monkeypatch.setattr(store, "CONFIG", str(tmp_path / "config"))
    os.makedirs(store.CONFIG, exist_ok=True)
    store.write_json(os.path.join(store.CONFIG, "settings.json"), {**SETTINGS, "measurement": {}})
    pid = "mapping-test"
    boq = {"import_id": "import-1", "name": "client.csv", "source_file": "client.csv",
           "source_fingerprint": "fingerprint-1", "rows": [_row()]}
    lines = [_line(), _line("line-other", "02.02", "Reinforced concrete footing alternative", quantity=4.1)]
    store.write_json(store.pdir(pid, "active_boq.json"), {"import_id": boq["import_id"]})
    store.write_json(store.pdir(pid, "boq", "import-1.json"), boq)
    store.write_json(store.pdir(pid, "runs", "run-1", "results.json"),
                     {"boq": lines, "elements": [{"uid": "F3-1", "cross_checks": [], "sources": []}]})
    os.makedirs(store.pdir(pid), exist_ok=True)
    with open(store.pdir(pid, "latest_run.txt"), "w", encoding="utf-8") as fh:
        fh.write("run-1")
    return pid, boq, lines


def _preview_and_decide(pid, action="approve", target="line-footing", version=0, comment="checked"):
    source_row = 8
    preview = mappings.preview(pid, source_row, action, target if action == "approve" else None, version)
    result = mappings.decide(pid, source_row, action, target if action == "approve" else None,
                             version, preview["preview_token"], comment)
    return preview, result


def test_approval_persists_and_is_reused_only_as_explicit_approval(mapping_project):
    pid, _, lines = mapping_project
    preview, decision = _preview_and_decide(pid)
    assert preview["changed"] and preview["after"]["target_key"] == "line-footing"
    assert decision["status"] == "approved" and decision["version"] == 1
    assert decision["reviewer"] and decision["timestamp"] and decision["comment"] == "checked"
    reloaded = mappings.load(pid)
    assert reloaded[decision["mapping_id"]]["history"][0]["action"] == "approve"
    rows = [_row()]
    applied, issues = mappings.applicable(pid, store.active_boq(pid), rows, lines)
    result = compare(rows, lines, {"F3-1": {"cross_checks": [], "sources": []}}, SETTINGS, applied, issues)["rows"][0]
    assert not issues and result["match_class"] == "approved"
    assert result["mapping_version"] == 1 and result["drawing_quantity"] == 3.6


def test_rejection_and_unresolved_decisions_are_persisted_without_matching(mapping_project):
    pid, boq, lines = mapping_project
    _preview_and_decide(pid, action="reject", target=None)
    rows = [dict(boq["rows"][0])]
    applied, issues = mappings.applicable(pid, boq, rows, lines)
    result = compare(rows, lines, {"F3-1": {"cross_checks": [], "sources": []}}, SETTINGS, applied, issues)["rows"][0]
    assert mappings.load(pid)[rows[0]["mapping_key"]]["status"] == "rejected"
    assert result["match_class"] == "rejected" and result["drawing_quantity"] is None
    preview = mappings.preview(pid, 8, "unresolved", expected_version=1)
    unresolved = mappings.decide(pid, 8, "unresolved", expected_version=1, preview_token=preview["preview_token"])
    assert unresolved["status"] == "unresolved" and unresolved["version"] == 2


def test_remapping_requires_new_preview_and_keeps_version_history(mapping_project):
    pid, _, _ = mapping_project
    _, first = _preview_and_decide(pid)
    preview = mappings.preview(pid, 8, "approve", "line-other", expected_version=first["version"])
    assert preview["before"]["drawing_quantity"] == 3.6
    assert preview["after"]["drawing_quantity"] == 4.1
    second = mappings.decide(pid, 8, "approve", "line-other", first["version"], preview["preview_token"])
    assert second["target_key"] == "line-other" and second["version"] == 2
    assert [h["target_key"] for h in second["history"]] == ["line-footing", "line-other"]
    with pytest.raises(ValueError, match="version changed"):
        mappings.decide(pid, 8, "approve", "line-footing", 1, preview["preview_token"])


def test_invalid_and_incompatible_targets_are_rejected(mapping_project):
    pid, _, lines = mapping_project
    with pytest.raises(ValueError, match="does not exist"):
        mappings.preview(pid, 8, "approve", "missing-target", 0)
    lines[0]["unit"] = "m2"
    store.write_json(store.run_path(pid, "run-1", "results.json"),
                     {"boq": lines, "elements": []})
    with pytest.raises(ValueError, match="Incompatible or missing units"):
        mappings.preview(pid, 8, "approve", "line-footing", 0)
    lines[0]["unit"] = "m3"
    lines[0]["category"] = "Beams"
    store.write_json(store.run_path(pid, "run-1", "results.json"),
                     {"boq": lines, "elements": []})
    with pytest.raises(ValueError, match="incompatible"):
        mappings.preview(pid, 8, "approve", "line-footing", 0)


def test_stale_source_and_missing_target_are_visible_in_review_queue(mapping_project):
    pid, boq, lines = mapping_project
    _, decision = _preview_and_decide(pid)
    boq["rows"][0]["description"] = "Changed source description"
    store.write_json(store.pdir(pid, "boq", "import-1.json"), boq)
    rows = [dict(boq["rows"][0])]
    applied, issues = mappings.applicable(pid, boq, rows, lines)
    assert decision["mapping_id"] in issues
    comparison = compare(rows, lines, {"F3-1": {"cross_checks": [], "sources": []}}, SETTINGS, applied, issues)
    assert "STALE_BOQ_MAPPING" in comparison["rows"][0]["flags"]
    queue = build_queue([], comparison, {}, {})
    assert any(item["code"] == "STALE_BOQ_MAPPING" for item in queue)

    lines.clear()
    rows = [dict(boq["rows"][0])]
    applied, issues = mappings.applicable(pid, boq, rows, lines)
    assert any("no longer present" in reason for reason in issues[decision["mapping_id"]])


def test_duplicate_approvals_are_blocked_and_history_does_not_rewrite_run(mapping_project):
    pid, boq, _ = mapping_project
    before = store.read_json(store.run_path(pid, "run-1", "results.json"))
    _preview_and_decide(pid)
    boq["rows"].append(_row(source_row=9, item_no="A-02", description="Concrete footing other"))
    store.write_json(store.pdir(pid, "boq", "import-1.json"), boq)
    preview = mappings.preview(pid, 9, "approve", "line-footing", 0)
    with pytest.raises(ValueError, match="already approved"):
        mappings.decide(pid, 9, "approve", "line-footing", 0, preview["preview_token"])
    assert store.read_json(store.run_path(pid, "run-1", "results.json")) == before


def test_duplicate_decision_and_preview_tokens_cannot_bypass_version_check(mapping_project):
    pid, _, _ = mapping_project
    first_preview = mappings.preview(pid, 8, "approve", "line-footing", 0)
    mappings.decide(pid, 8, "approve", "line-footing", 0, first_preview["preview_token"])
    with pytest.raises(ValueError, match="version changed"):
        mappings.decide(pid, 8, "approve", "line-footing", 0, first_preview["preview_token"])
    with pytest.raises(ValueError, match="preview"):
        mappings.decide(pid, 8, "approve", "line-other", 1, first_preview["preview_token"])


def test_canonical_mapping_is_separate_from_measured_quantity_and_versioned(mapping_project):
    pid, boq, lines = mapping_project
    canonical = mappings.save_canonical_item({"action": "create", "code": "SG-FOOT-01",
        "description": "Reinforced concrete footing", "category": "Footings", "unit": "m3"})
    target_id = f"canonical:{canonical['canonical_id']}"
    preview = mappings.preview(pid, 8, "approve", target_id, 0, "canonical")
    decision = mappings.decide(pid, 8, "approve", target_id, 0, preview["preview_token"],
                               "Canonical classification reviewed", "canonical")
    assert decision["target_type"] == "canonical" and decision["target_version"] == 1
    assert decision["target_code"] == "SG-FOOT-01"
    rows = [dict(boq["rows"][0])]
    applied, issues = mappings.applicable(pid, boq, rows, lines)
    comparison = compare(rows, lines, {"F3-1": {"cross_checks": [], "sources": []}},
                         SETTINGS, applied, issues)["rows"][0]
    assert comparison["mapping_status"] == "approved"
    assert comparison["mapping_target_type"] == "canonical"
    assert comparison["drawing_quantity"] is None
    assert comparison["match_class"] == "suggested"
    assert any("not measured quantities" in r for r in comparison["reasons"])
    library = mappings.canonical_library()
    assert library["version"] == 1 and len(library["history"]) == 1


def test_canonical_items_reject_incompatible_and_inactive_targets(mapping_project):
    pid, _, _ = mapping_project
    item = mappings.save_canonical_item({"action": "create", "code": "SG-BEAM-01",
        "description": "Concrete beam", "category": "Beams", "unit": "m3"})
    target = f"canonical:{item['canonical_id']}"
    with pytest.raises(ValueError, match="incompatible"):
        mappings.preview(pid, 8, "approve", target, 0, "canonical")
    inactive = mappings.save_canonical_item({"action": "deactivate", "canonical_id": item["canonical_id"],
                                              "expected_version": 1})
    assert inactive["status"] == "inactive" and inactive["version"] == 2
    with pytest.raises(ValueError, match="inactive"):
        mappings.preview(pid, 8, "approve", target, 0, "canonical")


def test_reject_can_record_an_incompatible_candidate_without_approving_it(mapping_project):
    pid, _, _ = mapping_project
    preview = mappings.preview(pid, 8, "reject", "line-other", 0, "calculated")
    rejected = mappings.decide(pid, 8, "reject", "line-other", 0, preview["preview_token"],
                               "Unit/category mismatch", "calculated")
    assert rejected["status"] == "rejected" and rejected["target_id"] == "line-other"


@pytest.mark.parametrize("mutation", ["description", "fingerprint", "import", "run", "target", "target_state", "action"])
def test_preview_is_bound_to_source_import_run_action_and_target(mapping_project, mutation):
    pid, boq, lines = mapping_project
    preview = mappings.preview(pid, 8, "approve", "line-footing", 0)
    if mutation == "description":
        boq["rows"][0]["description"] = "Altered description"
        store.write_json(store.pdir(pid, "boq", "import-1.json"), boq)
    elif mutation == "fingerprint":
        boq["source_fingerprint"] = "changed"
        store.write_json(store.pdir(pid, "boq", "import-1.json"), boq)
    elif mutation == "import":
        boq["import_id"] = "import-2"
        store.write_json(store.pdir(pid, "boq", "import-2.json"), boq)
        store.write_json(store.pdir(pid, "active_boq.json"), {"import_id": "import-2"})
    elif mutation == "run":
        os.makedirs(store.pdir(pid, "runs", "run-2"), exist_ok=True)
        store.write_json(store.run_path(pid, "run-2", "results.json"), {"boq": lines, "elements": []})
        with open(store.pdir(pid, "latest_run.txt"), "w", encoding="utf-8") as fh:
            fh.write("run-2")
    elif mutation == "target_state":
        result = store.read_json(store.run_path(pid, "run-1", "results.json"))
        result["boq"][0]["quantity"] = 9.9
        store.write_json(store.run_path(pid, "run-1", "results.json"), result)
    if mutation in ("target", "action"):
        target, action = ("line-other", "approve") if mutation == "target" else ("line-footing", "reject")
        with pytest.raises(ValueError, match="stale"):
            mappings.decide(pid, 8, action, target if action == "approve" else None, 0,
                            preview["preview_token"])
    else:
        with pytest.raises(ValueError, match="stale"):
            mappings.decide(pid, 8, "approve", "line-footing", 0, preview["preview_token"])


def test_preview_expiration_and_signature_tampering_are_rejected(mapping_project):
    pid, _, _ = mapping_project
    preview = mappings.preview(pid, 8, "approve", "line-footing", 0)
    claims, _ = preview["preview_token"].split(".", 1)
    payload = json.loads(base64.urlsafe_b64decode(claims + "=" * (-len(claims) % 4)))
    payload["issued_at"] -= mappings.PREVIEW_TTL_SECONDS + 1
    expired_claims = base64.urlsafe_b64encode(mappings._stable(payload).encode()).decode().rstrip("=")
    signature = hmac.new(mappings._preview_secret(), expired_claims.encode(), hashlib.sha256).hexdigest()
    with pytest.raises(ValueError, match="expired"):
        mappings.decide(pid, 8, "approve", "line-footing", 0, f"{expired_claims}.{signature}")
    tampered = preview["preview_token"][:-1] + ("0" if preview["preview_token"][-1] != "0" else "1")
    with pytest.raises(ValueError, match="modified"):
        mappings.decide(pid, 8, "approve", "line-footing", 0, tampered)


def test_corrupt_mapping_store_fails_closed_and_recovers_from_valid_backup(mapping_project):
    pid, _, _ = mapping_project
    _, first = _preview_and_decide(pid)
    preview = mappings.preview(pid, 8, "approve", "line-other", first["version"])
    mappings.decide(pid, 8, "approve", "line-other", first["version"], preview["preview_token"])
    path = store.pdir(pid, "boq", "mappings.json")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("{corrupt")
    with pytest.raises(ValueError, match="damaged"):
        mappings.load(pid)
    restored = mappings.recover(pid)
    assert next(iter(restored.values()))["version"] == 1
    assert list(__import__("pathlib").Path(os.path.dirname(path)).glob("mappings.json.corrupt.*"))
    with open(path, "w", encoding="utf-8") as fh:
        json.dump({"schema_version": 99, "mappings": {}}, fh)
    with pytest.raises(ValueError, match="Unsupported"):
        mappings.load(pid)


def test_wrong_project_mapping_record_is_not_reused(mapping_project):
    pid, boq, lines = mapping_project
    _, decision = _preview_and_decide(pid)
    foreign = dict(boq)
    key = decision["mapping_id"]
    rows = [dict(foreign["rows"][0])]
    # A record copied into another project's sidecar is rejected by its stored project identity.
    other_pid = "different-project"
    store.write_json(store.pdir(other_pid, "boq", "import-1.json"), foreign)
    store.write_json(store.pdir(other_pid, "boq", "mappings.json"),
                     {"schema_version": 1, "mappings": {key: decision}})
    _, issues = mappings.applicable(other_pid, foreign, rows, lines)
    assert issues[key] == ["Mapping belongs to a different project"]


def test_canonical_library_atomic_backup_and_corruption_recovery(mapping_project):
    mappings.save_canonical_item({"action": "create", "code": "SG-01", "description": "Footing",
                                  "category": "Footings", "unit": "m3"})
    mappings.save_canonical_item({"action": "create", "code": "SG-02", "description": "Second footing",
                                  "category": "Footings", "unit": "m3"})
    path = os.path.join(store.CONFIG, "canonical_boq_items.json")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("not json")
    with pytest.raises(ValueError, match="damaged"):
        mappings.canonical_library()
    restored = mappings.recover_canonical_library()
    assert restored["version"] == 1 and len(restored["items"]) == 1


def test_atomic_write_failure_keeps_last_valid_mapping(mapping_project, monkeypatch):
    pid, _, _ = mapping_project
    _, first = _preview_and_decide(pid)
    path = store.pdir(pid, "boq", "mappings.json")
    before = store.read_json(path)
    preview = mappings.preview(pid, 8, "approve", "line-other", first["version"])
    original_replace = mappings.os.replace

    def fail_current_file(source, destination):
        if destination == path:
            raise OSError("simulated atomic replace failure")
        return original_replace(source, destination)

    monkeypatch.setattr(mappings.os, "replace", fail_current_file)
    with pytest.raises(OSError, match="simulated"):
        mappings.decide(pid, 8, "approve", "line-other", first["version"], preview["preview_token"])
    assert store.read_json(path) == before
    assert store.read_json(path + ".bak") == before


def test_comparison_and_review_exports_include_mapping_identity_and_history(mapping_project, monkeypatch, tmp_path):
    from qto import service
    from qto.export import excel

    pid, boq, lines = mapping_project
    _, decision = _preview_and_decide(pid)
    boq["rows"][0]["description"] = "Changed description requires review"
    store.write_json(store.pdir(pid, "boq", "import-1.json"), boq)
    rows = [dict(boq["rows"][0])]
    applied, issues = mappings.applicable(pid, boq, rows, lines)
    comparison = compare(rows, lines, {"F3-1": {"cross_checks": [], "sources": []}},
                         SETTINGS, applied, issues)
    assert comparison["rows"][0]["mapping_id"] == decision["mapping_id"]
    comparison["boq"] = {"name": "synthetic.csv", "imported": "2025-01-01", "is_test": True}
    comparison["rows"][0]["boq_source_file"] = boq["source_file"]
    comparison["rows"][0]["boq_source_import_id"] = boq["import_id"]
    review = build_queue([], comparison, {}, {})
    monkeypatch.setattr(service, "results", lambda _: {"run": {"run_id": "run-1", "kind": "test",
        "finished": "2025-01-01", "started": "2025-01-01", "status": "complete", "counts": {}},
        "results": {"elements": [], "orphans": {}}})
    monkeypatch.setattr(service, "comparison", lambda _: comparison)
    monkeypatch.setattr(service, "queue", lambda _: review)
    monkeypatch.setattr(store, "project", lambda _: {"name": "Fixture", "location": "Test", "source_folder": "fixture"})
    monkeypatch.setattr(store, "settings", lambda _: SETTINGS)
    variance_path = str(tmp_path / "variance.xlsx")
    review_path = str(tmp_path / "review.xlsx")
    excel.export_variance(pid, variance_path)
    excel.export_review(pid, review_path)
    from openpyxl import load_workbook
    variance_sheet = load_workbook(variance_path, read_only=True)["Variance"]
    variance_values = list(variance_sheet.values)
    header_index = next(i for i, values in enumerate(variance_values) if "Source import ID" in values)
    headers = variance_values[header_index]
    row = variance_values[header_index + 1]
    assert "Source import ID" in headers and "Decision history" in headers
    assert decision["mapping_id"] in row and any("v1 approve" in str(value) for value in row)
    review_sheet = load_workbook(review_path, read_only=True)["Review Queue"]
    review_values = list(review_sheet.values)
    review_header_index = next(i for i, values in enumerate(review_values) if "Mapping history" in values)
    review_rows = review_values[review_header_index + 1:]
    mapping_row = next(values for values in review_rows if decision["mapping_id"] in values)
    assert any("v1 approve" in str(value) for value in mapping_row)


def _concurrent_decision(data_root, pid, source_row, target_id, token, output):
    from qto import store as child_store
    from qto.boq import mappings as child_mappings
    child_store.DATA = data_root
    try:
        child_mappings.decide(pid, source_row, "approve", target_id, 0, token)
        output.put("saved")
    except ValueError as exc:
        output.put(str(exc))


def test_cross_process_updates_serialize_and_stale_the_losing_preview(mapping_project):
    pid, boq, _ = mapping_project
    boq["rows"].append(_row(source_row=9, item_no="A-02", description="Other footing"))
    store.write_json(store.pdir(pid, "boq", "import-1.json"), boq)
    p1 = mappings.preview(pid, 8, "approve", "line-footing", 0)
    p2 = mappings.preview(pid, 9, "approve", "line-other", 0)
    ctx = multiprocessing.get_context("spawn")
    output = ctx.Queue()
    processes = [
        ctx.Process(target=_concurrent_decision, args=(store.DATA, pid, 8, "line-footing", p1["preview_token"], output)),
        ctx.Process(target=_concurrent_decision, args=(store.DATA, pid, 9, "line-other", p2["preview_token"], output)),
    ]
    for process in processes:
        process.start()
    for process in processes:
        process.join(20)
        assert process.exitcode == 0
    answers = [output.get(timeout=2), output.get(timeout=2)]
    assert answers.count("saved") == 1
    assert any("stale" in answer for answer in answers)
    assert len(mappings.load(pid)) == 1


def test_api_mapping_routes_enforce_project_scope_and_execute_preview_apply(mapping_project, monkeypatch):
    from fastapi import HTTPException
    from server import app as server_module
    pid, _, _ = mapping_project
    monkeypatch.setattr(store, "project", lambda project_id: {"id": pid} if project_id == pid else (_ for _ in ()).throw(KeyError(project_id)))
    from qto import service
    service.invalidate()
    with pytest.raises(HTTPException) as error:
        server_module.boq_mapping_preview("other-project", {"source_row": 8, "action": "approve"})
    assert error.value.status_code == 404
    preview = server_module.boq_mapping_preview(pid, {"source_row": 8, "action": "approve",
        "target_id": "line-footing", "target_type": "calculated", "expected_version": 0})
    saved = server_module.boq_mapping_decide(pid, {"source_row": 8, "action": "approve",
        "target_id": "line-footing", "target_type": "calculated", "expected_version": 0,
        "preview_token": preview["preview_token"], "comment": "API test"})
    assert saved["status"] == "approved"
    assert server_module.boq_mappings(pid)["mappings"][0]["mapping_id"] == saved["mapping_id"]
