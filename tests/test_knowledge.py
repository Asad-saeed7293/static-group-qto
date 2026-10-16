"""Tests for the knowledge layer: BOQ reading, item matching, benchmark classification, rules inheritance."""
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from qto.extract.tables import table_to_records_by_header  # noqa: E402
from qto.knowledge.benchmark import classify, run_benchmark  # noqa: E402
from qto.knowledge.boq_reader import read_static_boq  # noqa: E402
from qto.knowledge.model import load_model, match_row  # noqa: E402
from qto.store import load_rules_file  # noqa: E402


@pytest.fixture(scope="module")
def model():
    return load_model()


def _row(desc, sections=(), sheet="BOQ FOR VILLA"):
    return {"description": desc, "sheet": sheet, "section_names": list(sections)}


@pytest.mark.parametrize("desc,sections,expected", [
    ("RC Concrete work in Footings", ["Sub-Structure Works", "Villa"], "SUB-FTG"),
    ("Blinding Concrete", ["Sub-Structure Works", "Villa"], "SUB-BLD"),
    ("100mm thk Blinding Concrete", ["Sub-Structure Works", "Tanks"], "TANK-BLD"),
    ("RC Concrete works in Columns & Walls (GF + FF+ PH)", ["Super-Structural works"], "SUP-COLW"),
    ("Removal of excavated materials from the site", ["EARTH WORKS"], "EW-REM"),
    ("Bulk excavation in foundation", ["BOUNDARY WALL - (166.5 m/L x 3.15 m/H)", "Earth Works"], "BW-EXC"),
    ("200mm thk Block work", ["BOUNDARY WALL", "Masonry work"], "BW-BLK"),
    ("200mm thk Block wall", ["ARCHITECTURAL WORKS", "Internal Block works"], "BLK-INT"),
    ("350 mm thick double hollow block insulated wall", ["ARCHITECTURAL WORKS", "External Block works"], "BLK-EXT"),
    ("Electrical manhole (Size: 600mm/L x 600mm/W x 800mm/H)", ["Services"], "SRV-MH-EL"),
    ("External Wall plastering works in Boundarywall", ["External wall platering Works"], "PL-EXT-BW"),
])
def test_matching_uses_wording_and_section_context(model, desc, sections, expected):
    it, _ = match_row(_row(desc, sections), model)
    assert it is not None and it["id"] == expected


def test_classification_bands():
    assert classify(100.4, 100)[0] == "Exact Match"
    assert classify(103, 100)[0] == "Within Tolerance"
    assert classify(110, 100) == ("Difference", 10.0)
    assert classify(None, 100)[0] == "Unable to Calculate"
    assert classify(5, None)[0] == "Not Measured"


def test_read_static_boq_template(tmp_path):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Summary"
    s["B2"], s["C2"] = "Project:", "Test villa"
    s["B3"], s["C3"] = "Client:", "Someone"
    b = wb.create_sheet("BOQ FOR VILLA")
    b.append([None, None, "ITEM", "DESCRIPTION", "QTY", "UNIT", "RATE", "AMOUNT (QR)"])
    b.append([None, None, "b.1", "EARTH WORKS"])
    b.append([None, None, "1", "Bulk excavation in Isolated foudations, all tanks", 100.5, "m3"])
    b.append([None, None, "2", "Removal of excavated materials from the site", 100.5, "m3"])
    b.append([None, None, "4", "Anti-termite treatment", 1, "item", "ls"])
    b.append([None, None, "TOTAL CARRIED TO SUMMARY"])
    p = tmp_path / "boq.xlsx"
    wb.save(p)
    r = read_static_boq(str(p))
    assert r["meta"]["project"] == "Test villa"
    assert [x["kind"] for x in r["rows"]] == ["measured", "measured", "lump_sum"]
    assert r["rows"][0]["section_path"] == ["b.1"] and r["rows"][0]["quantity"] == 100.5


def test_benchmark_compares_engine_with_boq(model):
    boq = {"rows": [
        {"sheet": "BOQ FOR VILLA", "row": 1, "item": "1.2", "description": "RC Concrete work in Footings", "unit": "m3",
         "quantity": 10.0, "kind": "measured", "section_names": ["Sub-Structure Works"], "section_path": []},
        {"sheet": "BOQ FOR VILLA", "row": 2, "item": "1", "description": "Bulk excavation in Isolated foudations, all tanks",
         "unit": "m3", "quantity": 50.0, "kind": "measured", "section_names": ["EARTH WORKS"], "section_path": []}]}
    results = {"elements": [{"uid": "F1", "element_type": "footing", "floor": "FDN"}],
               "quantities": [{"element_uid": "F1", "quantity_type": "concrete_volume", "value": 10.2}]}
    b = run_benchmark(boq, results, model)
    assert b["rows"][0]["status"] == "Within Tolerance" and b["rows"][0]["engine_qty"] == 10.2
    assert b["rows"][1]["status"] == "Unable to Calculate"


def test_rules_inheritance_adds_project_differences():
    r = load_rules_file("rules_villa_std.json")
    assert r["floor_aliases"]["S.O.G"] == "GR"
    assert any(x["role"] == "column_outline" and "A-COLS-OTLN" in x["layers"] for x in r["layer_roles"])
    assert any(x["role"] == "footing_outline" for x in r["layer_roles"])          # inherited from the base file
    assert not any(f["id"] == "BS" for f in r["floors"])                          # no basement in these villas


def test_table_header_mapping_survives_offset_cells():
    t = {"char_height": 0.05, "rows": [
        [{"text": "FOOTING NO.", "x": 0.18}],
        [{"text": "W (mm)", "x": 0.54}, {"text": "L (mm)", "x": 0.91}, {"text": "D (mm)", "x": 1.27}],
        [{"text": "F17", "x": 0.18}, {"text": "2850", "x": 0.50}, {"text": "3250", "x": 0.93}, {"text": "550", "x": 1.27}, {"text": "T16/150", "x": 1.6}]]}
    rec = table_to_records_by_header(t, {"mark": "FOOTING NO.", "w_mm": "W (mm)", "l_mm": "L (mm)", "d_mm": "D (mm)"})[0]
    assert (rec["mark"], rec["w_mm"], rec["l_mm"], rec["d_mm"]) == ("F17", "2850", "3250", "550")
