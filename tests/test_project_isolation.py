import json
import os
import shutil
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from qto import service, store
from qto.changes import compare_runs
from qto.pipeline import Progress, run_full
from fastapi import HTTPException
from server.app import audit, project_items


@pytest.fixture
def isolated_projects(tmp_path, monkeypatch):
    config = tmp_path / "config"
    data = tmp_path / "data"
    config.mkdir()
    shutil.copyfile(os.path.join(store.CONFIG, "rules_villa_std.json"), config / "rules_villa_std.json")
    shutil.copyfile(os.path.join(store.CONFIG, "rules_ramzi_villa.json"), config / "rules_ramzi_villa.json")
    shutil.copyfile(os.path.join(store.CONFIG, "settings.json"), config / "settings.json")
    (config / "projects.json").write_text(json.dumps({"projects": [
        {"id": "isolation-a", "name": "Project A", "source_folder": str(tmp_path / "a-drawings"), "rules_file": "rules_villa_std.json"},
        {"id": "isolation-b", "name": "Project B", "source_folder": str(tmp_path / "b-drawings"), "rules_file": "rules_villa_std.json"},
    ]}), encoding="utf-8")
    monkeypatch.setattr(store, "CONFIG", str(config))
    monkeypatch.setattr(store, "DATA", str(data))
    service.invalidate()
    return tmp_path, data


def _write_run(pid, run_id, project_owner, results=None):
    store.write_json(store.run_path(pid, run_id, "run.json"), {"run_id": run_id, "project": project_owner, "status": "completed", "counts": {}})
    store.write_json(store.run_path(pid, run_id, "results.json"), results or {"elements": [], "quantities": [], "items": []})
    store.write_json(store.run_path(pid, run_id, "drawings.json"), [])


def test_project_without_run_does_not_read_another_projects_latest_run(isolated_projects):
    _write_run("isolation-b", "run-b", "isolation-b", {"items": [{"item_id": "B_ONLY"}]})
    store.set_latest("isolation-b", "run-b")

    assert service.results("isolation-a") is None


def test_revision_comparison_rejects_run_stored_under_another_project(isolated_projects):
    _write_run("isolation-a", "run-a", "isolation-b")

    with pytest.raises(KeyError, match="does not belong"):
        compare_runs("isolation-a", "run-a", "run-a")


def test_empty_processing_does_not_replace_latest_successful_run(isolated_projects):
    root, _ = isolated_projects
    drawings = root / "empty-drawings"
    drawings.mkdir()
    p = store.project("isolation-a")
    p["source_folder"] = str(drawings)
    config = json.loads((root / "config" / "projects.json").read_text(encoding="utf-8"))
    config["projects"][0]["source_folder"] = str(drawings)
    (root / "config" / "projects.json").write_text(json.dumps(config), encoding="utf-8")
    _write_run("isolation-a", "good-run", "isolation-a", {"items": [{"item_id": "A_ONLY"}]})
    store.set_latest("isolation-a", "good-run")

    progress = Progress()
    new_id = run_full("isolation-a", progress)

    assert store.latest_run_id("isolation-a") == "good-run"
    assert store.read_json(store.run_path("isolation-a", new_id, "run.json"))["status"] == "no drawings"
    assert progress.state["error"]


def test_synthetic_dxf_runs_through_real_pipeline_with_cad_traceability(isolated_projects):
    import ezdxf

    root, _ = isolated_projects
    drawings = root / "a-drawings"
    drawings.mkdir()
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 6
    doc.layers.new("S-FNDN-FTNG")
    doc.modelspace().add_lwpolyline(
        [(0, 0), (2, 0), (2, 3), (0, 3)],
        close=True,
        dxfattribs={"layer": "S-FNDN-FTNG"},
    )
    doc.saveas(str(drawings / "S-FP.dxf"))

    progress = Progress()
    run_id = run_full("isolation-a", progress)
    meta = store.read_json(store.run_path("isolation-a", run_id, "run.json"))
    result = store.read_project_run("isolation-a", run_id, "results.json")

    assert meta["status"] == "completed"
    assert store.latest_run_id("isolation-a") == run_id
    footing = next(e for e in result["elements"] if e["element_type"] == "footing")
    assert footing["dims"]["plan_area"]["value"] == 6.0
    assert footing["sources"][0]["handle"]
    assert footing["sources"][0]["drawing"] == "S-FP.dxf"
    assert service.results("isolation-b") is None

    from openpyxl import load_workbook
    from qto.export.excel import export_boq
    export_path = root / "project-a-boq.xlsx"
    export_boq("isolation-a", str(export_path))
    wb = load_workbook(export_path, read_only=True, data_only=False)
    detail_values = [value for row in wb["Calculation Detail"].iter_rows(values_only=True) for value in row]
    wb.close()
    assert "S-FP.dxf" in detail_values
    assert footing["sources"][0]["handle"] in detail_values


def test_unsupported_dxf_is_inventoried_without_becoming_a_successful_latest_run(isolated_projects):
    import ezdxf

    root, _ = isolated_projects
    drawings = root / "a-drawings"
    drawings.mkdir()
    doc = ezdxf.new("R2018")
    doc.modelspace().add_line((0, 0), (1, 1))
    doc.saveas(str(drawings / "unrecognized.dxf"))
    config = json.loads((root / "config" / "projects.json").read_text(encoding="utf-8"))
    config["projects"][0]["source_folder"] = str(drawings)
    (root / "config" / "projects.json").write_text(json.dumps(config), encoding="utf-8")
    _write_run("isolation-a", "good-run", "isolation-a", {"items": []})
    store.set_latest("isolation-a", "good-run")

    progress = Progress()
    run_id = run_full("isolation-a", progress)

    assert store.latest_run_id("isolation-a") == "good-run"
    assert store.read_json(store.run_path("isolation-a", run_id, "run.json"))["status"] == "unsupported drawings"
    assert service.results("isolation-a")["run"]["run_id"] == "good-run"


def test_project_items_endpoint_reads_only_the_requested_project(isolated_projects):
    _write_run("isolation-a", "run-a", "isolation-a", {"items": [{"item_id": "A_ONLY"}], "elements": [], "quantities": []})
    _write_run("isolation-b", "run-b", "isolation-b", {"items": [{"item_id": "B_ONLY"}], "elements": [], "quantities": []})
    store.set_latest("isolation-a", "run-a")
    store.set_latest("isolation-b", "run-b")

    assert [item["item_id"] for item in project_items("isolation-a")] == ["A_ONLY"]
    assert [item["item_id"] for item in project_items("isolation-b")] == ["B_ONLY"]


def test_unknown_project_cannot_write_or_read_an_audit_log(isolated_projects):
    with pytest.raises(HTTPException) as exc:
        audit("not-registered")

    assert exc.value.status_code == 404


def test_project_and_run_path_components_reject_traversal(isolated_projects):
    with pytest.raises(ValueError, match="Invalid project id"):
        store.pdir("..\\outside")
    with pytest.raises(ValueError, match="Invalid run id"):
        store.run_path("isolation-a", "..\\outside", "run.json")
