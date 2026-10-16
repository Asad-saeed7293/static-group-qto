"""Phase-3 tests: methodology rules, item calculators, Static Group BOQ, benchmark, new project, AI data tools.

Synthetic inputs are used where a calculator is tested in isolation; real-project checks use the latest
stored processing runs and are skipped when a project has not been processed.
"""
import copy
import json
import os
import shutil
import sys

import pytest
from shapely.geometry import box

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from qto import service, store  # noqa: E402
from qto.boq.static_template import DISPLAY, detect_new_items, generate_static_boq  # noqa: E402
from qto.changes import compare_results  # noqa: E402
from qto.extract.architecture import thickness_length  # noqa: E402
from qto.knowledge import rules_engine  # noqa: E402
from qto.knowledge.model import load_model  # noqa: E402
from qto.knowledge.rules_engine import STATUSES, effective, load_rules, settings_for  # noqa: E402
from qto.takeoff import items as I  # noqa: E402

KAWARI, RAMZI = "kawari-villa", "ramzi-villa"
need_kawari = pytest.mark.skipif(not store.latest_run_id(KAWARI), reason="Al-Kawari not processed")
need_ramzi = pytest.mark.skipif(not store.latest_run_id(RAMZI), reason="Ramzi not processed")
DECIDED = (I.AUTO, I.ASSUMPTION)


@pytest.fixture(scope="module")
def model():
    return load_model()


@pytest.fixture(scope="module")
def seed_rules():
    with open(rules_engine.SEED, encoding="utf-8") as fh:
        rules = {r["rule_id"]: dict(r, version=1, history=[], notes=[]) for r in json.load(fh)["rules"]}
    return rules


def _items(pid):
    return service.results(pid)["results"]["items"]


def _ctx(rules, inputs=None, model=None):
    return I.Ctx({"elements": [], "quantities": []}, {}, [], {}, rules, inputs or {}, model or load_model())


def _with_status(rules, rule_id, option):
    r = copy.deepcopy(rules)
    r[rule_id].update(status="CONFIRMED", selected=option)
    return r


# ------------------------------------------------------------------------- rule configuration
def test_rule_seed_is_consistent(seed_rules, model):
    assert len(seed_rules) >= 22
    for r in seed_rules.values():
        assert r["status"] in STATUSES, r["rule_id"]
        assert r["options"], f"{r['rule_id']} has no options"
        assert r.get("items"), f"{r['rule_id']} affects no item"
        for iid in r["items"]:
            assert iid in model["by_id"], f"{r['rule_id']} references unknown item {iid}"
        if r["status"] == "UNKNOWN":
            assert not r.get("selected"), f"UNKNOWN rule {r['rule_id']} must not have a selected option"
        else:
            assert r.get("selected") in r["options"], f"{r['rule_id']} selected option not among options"


def test_effective_unknown_gives_no_option(seed_rules):
    unknown = next(r for r in seed_rules.values() if r["status"] == "UNKNOWN")
    assert effective(seed_rules, unknown["rule_id"]) == (None, "UNKNOWN")
    assert effective(seed_rules, "no_such_rule") == (None, "UNKNOWN")
    opt = next(iter(unknown["options"]))
    assert effective(_with_status(seed_rules, unknown["rule_id"], opt), unknown["rule_id"]) == (opt, "CONFIRMED")


def test_settings_follow_approved_rule_only(seed_rules):
    base = store.base_settings()
    assert settings_for(seed_rules, base)["measurement"] == base["measurement"], "UNKNOWN rules must not change settings"
    r = seed_rules["column_height_method"]
    opt = next(k for k, v in r["variants"].items() if v.get("measurement"))
    s = settings_for(_with_status(seed_rules, "column_height_method", opt), base)
    for k, v in r["variants"][opt]["measurement"].items():
        assert s["measurement"][k] == v


@pytest.fixture
def isolated_rule_state(tmp_path, monkeypatch):
    monkeypatch.setattr(rules_engine, "STATE", str(tmp_path / "rule_state.json"))
    events = []
    monkeypatch.setattr(store, "audit", lambda pid, event, **d: events.append((pid, event, d)))
    return events


def test_rule_approval_versioning_and_audit(isolated_rule_state):
    rid = "lintel_bearing"
    r0 = load_rules()[rid]
    assert r0["status"] == "UNKNOWN" and r0["version"] == 1
    opt = next(iter(r0["options"]))
    r1 = rules_engine.update_rule(rid, "approve", user="estimator", option=opt, note="Per Static Group practice")
    assert (r1["status"], r1["selected"], r1["version"], r1["approved_by"]) == ("CONFIRMED", opt, 2, "estimator")
    assert r1["approved_at"] and r1["history"][-1]["before"]["status"] == "UNKNOWN"
    assert r1["notes"][-1]["text"] == "Per Static Group practice"
    r2 = rules_engine.update_rule(rid, "note", user="ceo", note="question")
    assert r2["version"] == 2 and len(r2["notes"]) == 2
    r3 = rules_engine.update_rule(rid, "reject", user="estimator")
    assert (r3["status"], r3["selected"], r3["version"]) == ("UNKNOWN", None, 3)
    assert [e[1] for e in isolated_rule_state] == ["rule_approve", "rule_note", "rule_reject"]
    assert all(e[0] == "platform" for e in isolated_rule_state)


def test_rule_approval_rejects_bad_input(isolated_rule_state):
    with pytest.raises(ValueError):
        rules_engine.update_rule("lintel_bearing", "approve", option="not-an-option")
    with pytest.raises(KeyError):
        rules_engine.update_rule("nope", "approve", option="x")
    with pytest.raises(ValueError):
        rules_engine.update_rule("lintel_bearing", "edit", status="MAYBE")


# ------------------------------------------------------------------------- unknown rules / missing data (all projects)
@pytest.mark.parametrize("pid", [RAMZI, KAWARI, "sulaiti-villa"])
def test_unknown_rule_never_reports_a_decided_quantity(pid):
    if not store.latest_run_id(pid):
        pytest.skip("not processed")
    for it in _items(pid):
        open_rules = [r for r in it["rules"] if r["selected"] is None]
        if open_rules:
            assert it["status"] not in DECIDED, f"{it['item_id']} {it['variant']} is {it['status']} with open rule(s)"
            if not it["item_id"].startswith(("TANK-", "WP-TANK", "BW-", "PL-EXT-BW")):   # those come from estimator-typed inputs
                assert it["value"] is None, f"{it['item_id']} {it['variant']} has a value although a rule is open"
        if it["status"] in (I.INPUT, I.NA_DRAWING) and it["item_id"] not in ("EW-DW", "BW-DW"):
            assert it["missing"], f"{it['item_id']} is {it['status']} without saying what is missing"


# ------------------------------------------------------------------------- walls / blockwork / plaster / lintels
def test_wall_thickness_from_hatch_polygon():
    t, L, snapped = thickness_length(box(0, 0, 5.0, 0.2))
    assert snapped and t == pytest.approx(0.2, abs=1e-6) and L == pytest.approx(5.0, abs=1e-6)


@pytest.mark.parametrize("t,label", [(0.345, 0.35), (0.198, 0.2), (0.26, 0.25), (0.27, 0.27), (0.2, 0.2), (0.62, 0.62)])
def test_thickness_label_snapping(t, label):
    assert I._thk_label(t) == pytest.approx(label)


@need_kawari
def test_blockwork_gross_candidates_and_no_value_while_unknown():
    blk = [i for i in _items(KAWARI) if i["item_id"] in ("BLK-EXT", "BLK-INT")]
    assert blk, "no blockwork items"
    rules = load_rules()
    for i in blk:
        assert i["variant"].endswith(("external", "internal"))
        assert any(k.startswith("gross") for k in i["candidates"]), i["candidates"]
        assert {r["rule_id"] for r in i["rules"]} >= {"wall_height_method", "blockwork_opening_deduction"}
        if rules["wall_height_method"]["status"] == "UNKNOWN":
            assert i["value"] is None
    assert {i["item_id"] for i in blk} == {"BLK-EXT", "BLK-INT"}


@need_kawari
def test_plaster_candidates_name_faces_and_heights():
    pl = {i["item_id"]: i for i in _items(KAWARI) if i["item_id"] in ("PL-EXT", "PL-INT")}
    assert set(pl) == {"PL-EXT", "PL-INT"}
    assert all("outer face" in k for k in pl["PL-EXT"]["candidates"])
    assert any("both faces internal" in k for k in pl["PL-INT"]["candidates"])
    assert any("Staircase" in m for m in pl["PL-INT"]["missing"])


@need_kawari
def test_lintel_flags_partial_opening_detection():
    li = next(i for i in _items(KAWARI) if i["item_id"] == "LINTEL")
    assert li["missing"] and "partial" in li["missing"][0]
    assert any("bearing" in k for k in li["candidates"])
    assert li["rules"][0]["rule_id"] == "lintel_bearing"


# ------------------------------------------------------------------------- earthworks / waterproofing / manholes
@need_kawari
def test_earthworks_require_ground_level_and_show_working_space_candidates():
    it = {i["item_id"]: i for i in _items(KAWARI)}
    inputs = store.project_inputs(KAWARI)
    if inputs.get("natural_ground_level") in (None, ""):
        assert it["EW-EXC"]["status"] == I.INPUT and it["EW-EXC"]["value"] is None
        assert any("ground level" in m.lower() for m in it["EW-EXC"]["missing"])
    ws = {k.split("working space ")[1].split(" ")[0] for k in it["EW-EXC"]["candidates"] if "working space" in k}
    assert {"0.0", "0.3", "0.5"} <= ws, "working space must be configurable candidates, not an assumption"
    assert it["EW-REM"]["rules"][0]["rule_id"] == "removal_equals_excavation"


@need_kawari
def test_waterproofing_only_candidates_from_drawn_elements():
    for iid in ("WP-SUB-H", "WP-SUB-V"):
        w = next(i for i in _items(KAWARI) if i["item_id"] == iid)
        assert set(w["candidates"]) == {"sides_and_tops", "sides_only"}
        assert w["status"] not in DECIDED or w["rules"][0]["selected"]


@need_kawari
def test_manholes_counted_from_drawing_symbols():
    mh = {i["item_id"]: i for i in _items(KAWARI) if i["item_id"].startswith("SRV-")}
    assert {"SRV-MH-DR", "SRV-MH-RW", "SRV-SOAK", "SRV-MH-EL"} <= set(mh)
    assert any("blocks matching" in k for k in mh["SRV-MH-DR"]["candidates"])
    assert mh["SRV-MH-EL"]["rules"][0]["rule_id"] == "electrical_manhole_source"


# ------------------------------------------------------------------------- tanks / boundary wall (estimator inputs)
def test_tanks_require_input_then_follow_formula(seed_rules, model):
    out = I.tank_items(_ctx(seed_rules, {}, model), model["by_id"])
    assert {r["status"] for r in out} == {I.INPUT} and all(r["value"] is None for r in out)
    tank = {"name": "T1", "L": 3, "W": 2, "H": 2, "wall_t": 0.2, "base_t": 0.3, "slab_t": 0.2}
    v = {r["item_id"]: r for r in I.tank_items(_ctx(seed_rules, {"tanks": [tank]}, model), model["by_id"])}
    assert v["TANK-FTG"]["value"] == pytest.approx(3.4 * 2.4 * 0.3, abs=1e-3)
    assert v["TANK-WALL"]["value"] == pytest.approx(2 * 5.4 * 0.2 * 2, abs=1e-3)
    assert v["WP-TANK-V"]["value"] == pytest.approx(2 * (3.4 + 2.4) * 2.5, abs=1e-3)
    assert v["TANK-FTG"]["status"] == I.REVIEW, "tank geometry source rule is UNKNOWN"
    partial = I.tank_items(_ctx(seed_rules, {"tanks": [{"name": "T2", "L": 3}]}, model), model["by_id"])
    assert all(r["status"] == I.INPUT and r["value"] is None and "T2" in r["missing"][0] for r in partial)


def test_boundary_wall_inputs(seed_rules, model):
    out = I.boundary_items(_ctx(seed_rules, {"boundary_wall": {"length": 100}}, model), model["by_id"])
    assert all(r["status"] == I.INPUT and "height" in r["missing"][0] for r in out)
    bw = dict(length=100, height=3, footing_w=0.8, footing_t=0.3, exc_depth=1.2, gb_w=0.2, gb_d=0.5, col_w=0.2, col_d=0.4,
              col_spacing=4, neck_h=0.6, stiff_w=0.2, stiff_d=0.2, block_t=0.2)
    v = {r["item_id"]: r["value"] for r in I.boundary_items(_ctx(seed_rules, {"boundary_wall": bw}, model), model["by_id"])}
    assert v["BW-EXC"] == pytest.approx(100 * 1.0 * 1.2)
    assert v["BW-REM"] == v["BW-EXC"]
    assert v["BW-BF"] == pytest.approx(v["BW-EXC"] - v["BW-BLD"] - v["BW-FTG"] - v["BW-GB"] - v["BW-NECK"], abs=1e-3)
    assert v["BW-COL"] == pytest.approx(26 * 0.2 * 0.4 * 3)


# ------------------------------------------------------------------------- Static Group BOQ template / mapping
@need_kawari
def test_static_boq_preserves_structure_and_statuses():
    b = generate_static_boq(_items(KAWARI))
    assert b["lines"]
    # "code" matches Static Group's own convention (plain 1, 2, 3 ... reset per section/subsection) and so
    # deliberately repeats across sections - the historical BOQs do the same (e.g. "1" appears in many
    # sections). "key" (item_id + variant) is the one field guaranteed unique across the whole BOQ.
    keys = [l["key"] for l in b["lines"]]
    assert len(keys) == len(set(keys)), "duplicate BOQ line keys"
    codes_per_sheet_section = [(l["sheet"], l["section"], l["subsection"], l["code"]) for l in b["lines"]]
    assert len(codes_per_sheet_section) == len(set(codes_per_sheet_section)), "duplicate code within the same section"
    for l in b["lines"]:
        for k in ("code", "key", "sheet", "section", "description", "unit", "calculation_status", "display_status", "source", "confidence", "review_status"):
            assert k in l
        assert l["sheet"]
        if l.get("is_group_header"):
            continue
        assert l["display_status"] in set(DISPLAY.values())
        if l["quantity"] is None:
            assert l["display_status"] not in ("Calculated", "Estimated")
        if l["rules"]:
            assert l["display_status"] not in ("Calculated", "Estimated")
    assert not [l for l in b["lines"] if l["section"] == "X"], "every calculated item should be placed in the template"


@need_kawari
def test_preview_matches_export_structure(tmp_path):
    """The in-app preview (preview_static_boq) must show exactly the sheets/sections/subtotals that
    Export Excel actually writes - otherwise someone could preview one layout and download another."""
    from qto.export.excel import export_static_boq, preview_static_boq

    pv = preview_static_boq(KAWARI)
    assert pv["sheet_order"] == ["Preliminaries", "BOQ FOR VILLA", "Plastering works"]
    assert pv["summary"], "no sections closed into the preview's Summary"

    out = tmp_path / "kawari_boq.xlsx"
    export_static_boq(KAWARI, str(out))
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    assert set(wb.sheetnames) >= {*pv["sheet_order"], "Summary", "Calculation Detail"}

    sm = wb["Summary"]
    exported_codes = [sm.cell(row=r, column=3).value for r in range(6, sm.max_row + 1) if sm.cell(row=r, column=3).value]
    exported_codes = [c for c in exported_codes if c != "GRAND TOTAL (QR)"]
    assert exported_codes == [s["code"] for s in pv["summary"]], "preview's Summary order must match the exported file's"

    for sheet_name in pv["sheet_order"]:
        preview_codes = [r["code"] for r in pv["sheets"][sheet_name] if r["row"] in ("section", "subsection")]
        ws = wb[sheet_name]
        written_codes = [ws.cell(row=r, column=3).value for r in range(6, ws.max_row + 1)
                         if isinstance(ws.cell(row=r, column=3).value, str) and any(c.isalpha() for c in ws.cell(row=r, column=3).value)
                         and ws.cell(row=r, column=4).value and ws.cell(row=r, column=3).font.bold]
        assert preview_codes == written_codes, f"{sheet_name}: preview section/subsection codes must match what gets exported"


@need_kawari
def test_export_changes(tmp_path):
    """The drawing-revision changes report must mirror compare_runs()'s own numbers exactly - it is a
    shareable export of that comparison, not a separate calculation."""
    from qto.changes import compare_runs
    from qto.export.excel import export_changes

    runs = sorted(r["run_id"] for r in store.list_runs(KAWARI))
    old, new = runs[0], runs[1]
    cmp = compare_runs(KAWARI, old, new)
    out = tmp_path / "changes.xlsx"
    export_changes(KAWARI, str(out), old, new)
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    assert set(wb.sheetnames) == {"Summary", "Drawings", "Elements", "BOQ Items"}
    ws2 = wb["Elements"]
    written_elements = ws2.max_row - 6  # minus title/header block rows (no blank-state path here)
    assert written_elements == len(cmp["elements"])
    ws4 = wb["BOQ Items"]
    assert ws4.max_row - 6 == len(cmp["items"])
    if cmp["items"]:
        assert ws4.cell(row=7, column=1).value == cmp["items"][0]["item_id"]


@need_kawari
def test_export_tender_pdf(tmp_path):
    """The tender PDF must be built from the exact same row layout as the preview/export, contain every
    sheet, and never silently drop description text (reportlab treats Paragraph text as light XML, so an
    unescaped '<' in a description - e.g. the '<L>' placeholder in some item descriptions - would vanish)."""
    from qto.export.excel import preview_static_boq
    from qto.export.pdf import export_tender_pdf

    pv = preview_static_boq(KAWARI)
    out = tmp_path / "tender.pdf"
    export_tender_pdf(KAWARI, str(out))
    assert out.exists() and out.stat().st_size > 1000

    from pypdf import PdfReader
    text = "\n".join(p.extract_text() for p in PdfReader(str(out)).pages)
    assert pv["project"]["name"] in text
    assert "<L>" in text, "angle-bracket placeholder text must render literally, not be swallowed as markup"
    for sheet in pv["sheet_order"]:
        assert sheet in text
    for s in pv["summary"]:
        assert s["code"] in text and s["title"] in text


def test_new_boq_item_detected():
    rows = [{"row": 1, "description": "RC Concrete work in Footings", "sheet": "BOQ", "section_names": ["Sub-Structure Works"]},
            {"row": 2, "description": "Supply and install rooftop solar photovoltaic array", "unit": "kWp", "quantity": 10, "section_names": []}]
    new = detect_new_items(rows)
    assert [n["row"] for n in new] == [2] and new[0]["status"] == "New BOQ item detected"


def test_unit_normalisation_in_model(model):
    from qto.boq.compare import norm_unit
    assert [norm_unit(u) for u in ("m³", "cum", "Sqm", "lm", "Nos")] == ["m3", "m3", "m2", "m", "nr"]
    for it in model["items"]:
        if it["unit"] == "-":
            assert it.get("kind") == "excluded", it["id"]
        else:
            assert norm_unit(it["unit"]) in {"m3", "m2", "m", "nr", "item"}, (it["id"], it["unit"])


# ------------------------------------------------------------------------- historical benchmark
@need_kawari
def test_benchmark_lines_have_rule_and_diagnosis():
    b = store.read_json(store.pdir(KAWARI, "benchmark.json"))
    assert b and b["rows"]
    measured = [r for r in b["rows"] if r["kind"] == "measured"]
    assert sum(b["summary"]["by_status"].values()) == len(measured)
    beam = next(r for r in measured if r.get("item_id") == "SUP-BEAM")
    assert beam["engine_qty"] is None and beam["status"] in ("Candidates Only", "Unable to Calculate")
    assert any("UNKNOWN" in c["evidence"] for c in beam["diagnosis"])
    assert any("not a conclusion" in c["cause"] + c["evidence"] for c in beam["diagnosis"])
    for r in measured:
        if r["status"] in ("Exact Match", "Within Tolerance", "Difference"):
            assert r["engine_qty"] is not None and r["diff_pct"] is not None


# ------------------------------------------------------------------------- change detection
def test_change_detection_between_runs():
    old = {"elements": [{"uid": "A", "element_type": "beam", "dims": {"w": 0.2}, "drawing": "x"}, {"uid": "B", "element_type": "beam", "dims": {}, "drawing": "x"}],
           "quantities": [{"element_uid": "A", "quantity_type": "concrete_volume", "value": 1.0}],
           "items": [{"item_id": "SUP-BEAM", "variant": "", "description": "Beams", "unit": "m3", "value": 10.0, "status": I.ASSUMPTION}]}
    new = {"elements": [{"uid": "A", "element_type": "beam", "dims": {"w": 0.3}, "drawing": "x"}, {"uid": "C", "element_type": "column", "dims": {}, "drawing": "y"}],
           "quantities": [{"element_uid": "A", "quantity_type": "concrete_volume", "value": 1.5}],
           "items": [{"item_id": "SUP-BEAM", "variant": "", "description": "Beams", "unit": "m3", "value": 12.0, "status": I.ASSUMPTION}]}
    d = compare_results(old, new, {"a.dxf": {"file": "a.dxf", "size": 1, "entities": 5, "extents": []}},
                        {"a.dxf": {"file": "a.dxf", "size": 2, "entities": 5, "extents": []}, "b.dxf": {"file": "b.dxf", "size": 1, "entities": 1, "extents": []}})
    s = d["summary"]
    assert (s["elements_added"], s["elements_removed"], s["elements_modified"], s["drawings_changed"]) == (1, 1, 1, 2)
    a = next(e for e in d["elements"] if e["uid"] == "A")
    assert a["dims"]["w"] == [0.2, 0.3] and a["quantities"]["concrete_volume"] == [1.0, 1.5]
    assert d["items"][0]["delta"] == 2.0 and d["items"][0]["delta_pct"] == 20.0
    assert (s["concrete_old_m3"], s["concrete_new_m3"]) == (1.0, 1.5)


# ------------------------------------------------------------------------- AI data retrieval
def test_ai_tools_cite_their_sources():
    from qto.ai import tools as T
    for name in ("list_projects", "unknown_rules", "concrete_all"):
        r = T.TOOLS[name][0]()
        assert "data" in r and r["citations"], name
    rows = T.unknown_rules()["data"]
    assert {r["rule_id"] for r in rows} == {r["rule_id"] for r in load_rules().values() if r["status"] != "CONFIRMED"}


@need_kawari
def test_ai_status_numbers_match_stored_items():
    from qto.ai.tools import project_status
    s = project_status(KAWARI)["data"]
    items = _items(KAWARI)
    assert s["boq_items"] == len(items)
    assert s["items_needing_review"] == sum(i["status"] == I.REVIEW for i in items)
    assert s["concrete_calculated_m3"] == pytest.approx(sum(i["value"] for i in items if i["value"] is not None and i["unit"] == "m3" and i["category"] == "Concrete"), abs=0.06)


@need_kawari
def test_assistant_explains_difference_without_concluding(monkeypatch):
    from qto.ai import agent
    monkeypatch.setattr(store, "base_settings", lambda: {"ai": {"enabled": False}})
    a = agent.answer("Why is beam concrete different from the historical BOQ?", KAWARI)
    beam = next(r for r in store.read_json(store.pdir(KAWARI, "benchmark.json"))["rows"] if r.get("item_id") == "SUP-BEAM")
    assert a["mode"] == "deterministic"
    assert f"{beam['historical_qty']:,.2f}" in a["answer"]
    assert "No conclusion has been made" in a["answer"]
    assert any(c["type"] == "benchmark" for c in a["citations"])


def test_assistant_refuses_to_invent(monkeypatch):
    from qto.ai import agent
    monkeypatch.setattr(store, "base_settings", lambda: {"ai": {"enabled": False}})
    a = agent.answer("What is the price of steel in Doha tomorrow?", None)
    assert a["citations"] == [] and "could not match" in a["answer"]


def test_assistant_matches_project_names():
    from qto.ai.agent import _projects_in_text
    ids = {p["id"] for p in store.projects()}
    if {"sulaiti-villa", "kawari-villa"} <= ids:
        assert _projects_in_text("Compare Sulaiti with Al-Kawari") == ["sulaiti-villa", "kawari-villa"]


def test_daily_report_is_data_only():
    from qto.reports.daily import build_daily_report
    rep = build_daily_report(save=False)
    assert len(rep["pending_approvals"]) == sum(r["status"] != "CONFIRMED" for r in load_rules().values())
    assert rep["not_measured"] and "Time saved" in rep["not_measured"][0]
    assert "# Static Group daily project report" in rep["text"]
    for p in rep["projects"]:
        if p["processed"]:
            assert p["items_complete"] <= p["boq_items"]


# ------------------------------------------------------------------------- new project workflow (end to end on real drawings)
@need_ramzi
def test_new_project_from_uploaded_drawings(tmp_path, monkeypatch):
    from qto.pipeline import Progress, run_full
    src = store.project(RAMZI)["source_folder"]
    files = {}
    for root, _, fs in os.walk(src):
        for f in fs:
            if f in ("S-FP.dxf", "S-FD.dxf") and "FOUNDATIONS" in root:
                files[f] = os.path.join(root, f)
    if len(files) != 2:
        pytest.skip("Ramzi foundation drawings not found")
    ftg_full = next(i for i in _items(RAMZI) if i["item_id"] == "SUB-FTG")
    cfg = tmp_path / "config"
    shutil.copytree(store.CONFIG, cfg)
    with open(cfg / "projects.json", "w", encoding="utf-8") as fh:
        json.dump({"projects": []}, fh)
    monkeypatch.setattr(store, "CONFIG", str(cfg))
    monkeypatch.setattr(store, "DATA", str(tmp_path / "data"))
    drawings = tmp_path / "upload"
    drawings.mkdir()
    for f, p in files.items():
        shutil.copy(p, drawings / f)
    store.add_project({"id": "new-villa", "name": "New Villa", "source_folder": str(drawings), "rules_file": "rules_ramzi_villa.json",
                       "historical_boq": None, "role": "new project"})
    with pytest.raises(ValueError):
        store.add_project({"id": "new-villa", "name": "dup"})
    run_id = run_full("new-villa", Progress())
    res = store.read_json(store.run_path("new-villa", run_id, "results.json"))
    ftg_new = next(i for i in res["items"] if i["item_id"] == "SUB-FTG")
    assert ftg_new["value"] == pytest.approx(ftg_full["value"], rel=1e-6), "footings need only the plan + schedule"
    sup = next(i for i in res["items"] if i["item_id"] == "SUP-SLAB")
    assert sup["value"] in (None, 0) or sup["status"] != I.AUTO
    b = generate_static_boq(res["items"])
    assert b["summary"].get("Missing", 0) + b["summary"].get("Needs Review", 0) > 0
    assert os.path.exists(store.pdir("new-villa", "audit_log.jsonl"))
