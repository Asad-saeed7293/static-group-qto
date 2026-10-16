"""Build docs/../Static_Group_QTO_Features_Explained.pptx - every feature of the system, in simple words.

Run:  python tools/make_feature_deck.py
Uses the official STATIC logo (web/brand) and real screenshots of the application (docs/ppt_assets).
All numbers are the verified figures of the current reference projects (2026-10-08).
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = os.path.join(ROOT, "docs", "ppt_assets")
LOGO_DARK = os.path.join(ROOT, "web", "brand", "static-logo.png")
LOGO_WHITE = os.path.join(ROOT, "web", "brand", "static-logo-white.png")
OUT = os.path.join(ROOT, "Static_Group_QTO_Features_Explained.pptx")

INK, GREY, LIGHT, LINE, WHITE = RGBColor(0x11, 0x12, 0x14), RGBColor(0x5F, 0x64, 0x6D), RGBColor(0xF4, 0xF5, 0xF6), RGBColor(0xE2, 0xE4, 0xE7), RGBColor(0xFF, 0xFF, 0xFF)
GREEN, BLUE, AMBER, RED, NEUTRAL = RGBColor(0x1D, 0x70, 0x48), RGBColor(0x2B, 0x5A, 0x8A), RGBColor(0xB7, 0x7A, 0x12), RGBColor(0xB4, 0x23, 0x18), RGBColor(0x8A, 0x91, 0x99)
FONT = "Segoe UI"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height
page = [0]


def txt(slide, x, y, w, h, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(size), color, bold, font
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius:
        s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    s.text_frame.margin_left = s.text_frame.margin_right = Inches(0.12)
    return s


def boxtext(s, lines, size=14, color=INK, bold_first=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size if i or not bold_first else size + 1), color
        r.font.bold = bool(bold_first and i == 0)
        p.space_after = Pt(3)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def base(title, kicker=""):
    page[0] += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    bar.fill.solid(); bar.fill.fore_color.rgb = INK; bar.line.fill.background()
    if kicker:
        txt(s, Inches(0.6), Inches(0.35), Inches(9), Inches(0.35), kicker.upper(), 11, GREY, True)
    txt(s, Inches(0.6), Inches(0.62), Inches(10.0), Inches(0.75), title, 28 if len(title) <= 46 else 23 if len(title) <= 58 else 20, INK, True)
    s.shapes.add_picture(LOGO_DARK, Inches(10.85), Inches(0.5), height=Inches(0.4))
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6), Inches(1.45), W - Inches(0.6), Inches(1.45))
    ln.line.color.rgb = LINE
    txt(s, Inches(0.6), H - Inches(0.45), Inches(8), Inches(0.3), "Static Group · Automated Quantity Takeoff & BOQ System", 9, GREY)
    txt(s, W - Inches(1.4), H - Inches(0.45), Inches(0.8), Inches(0.3), str(page[0]), 9, GREY, align=PP_ALIGN.RIGHT)
    return s


def bullets(s, items, x=Inches(0.6), y=Inches(1.75), w=Inches(12), size=18, gap=8):
    tb = s.shapes.add_textbox(x, y, w, H - y - Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        sub = it.startswith("  ")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ("–  " if sub else "•  ") + it.strip()
        r.font.name, r.font.size = FONT, Pt(size - 3 if sub else size)
        r.font.color.rgb = GREY if sub else INK
        p.level = 1 if sub else 0
        p.space_after = Pt(gap)
    return tb


def shot(s, name, x, y, w, maxh=None):
    p = os.path.join(A, name)
    if os.path.exists(p):
        maxh = maxh or (H - y - Inches(0.7))
        ratio = 0.625                                  # screenshots are 16:10
        if w * ratio > maxh:
            w = int(maxh / ratio)
        pic = s.shapes.add_picture(p, x, y, width=w)
        pic.line.color.rgb = LINE
        pic.line.width = Pt(1)
        return pic


def cards(s, items, cols=3, y=Inches(1.8), h=Inches(1.55), size=13, colors=None, x0=Inches(0.6), total_w=None):
    total_w = total_w or (W - Inches(1.2))
    gap = Inches(0.25)
    cw = int((total_w - gap * (cols - 1)) / cols)
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        b = box(s, x0 + c * (cw + gap), y + r * (h + gap), cw, h, fill=LIGHT if not colors else WHITE)
        if colors:
            stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0 + c * (cw + gap), y + r * (h + gap), Inches(0.09), h)
            stripe.fill.solid(); stripe.fill.fore_color.rgb = colors[i]; stripe.line.fill.background()
        boxtext(b, it, size)


def flow(s, steps, y=Inches(2.2), h=Inches(1.25), color=INK):
    n = len(steps)
    gap = Inches(0.35)
    bw = int((W - Inches(1.2) - gap * (n - 1)) / n)
    for i, (t, d) in enumerate(steps):
        x = Inches(0.6) + i * (bw + gap)
        b = box(s, x, y, bw, h, fill=color if i in (0, n - 1) else LIGHT, line=None)
        boxtext(b, [t, d], 12, WHITE if i in (0, n - 1) else INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < n - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Inches(0.06), y + h / 2 - Inches(0.12), gap - Inches(0.12), Inches(0.24))
            a.fill.solid(); a.fill.fore_color.rgb = NEUTRAL; a.line.fill.background()


# ===================================================================== slides
# 1 title
page[0] += 1
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = INK
s.shapes.add_picture(LOGO_WHITE, Inches(0.8), Inches(0.8), height=Inches(0.9))
txt(s, Inches(0.8), Inches(2.4), Inches(12.0), Inches(1.0), "Automated Quantity Takeoff & BOQ System", 38, WHITE, True)
txt(s, Inches(0.8), Inches(3.45), Inches(11.5), Inches(0.8), "Every feature explained in simple words", 24, RGBColor(0xC9, 0xCC, 0xD1))
txt(s, Inches(0.8), Inches(5.6), Inches(11.5), Inches(1.0), ["AI-assisted · the estimator stays in control · every number shows where it came from",
                                                              "Status: ready for a controlled live test with a real Static Group project · October 2026"], 14, RGBColor(0x9A, 0x9F, 0xA6))
notes(s, "This deck walks through every part of the system in plain language. The short version: the computer reads the drawings and proposes quantities; Static Group's estimators check, correct and approve them. Nothing becomes final without a person.")

# 2 the problem today
s = base("The problem we are solving", "Why")
cards(s, [["Manual take-off is slow", "An estimator measures every footing, column, beam, slab and wall by hand from dozens of drawings."],
          ["Hard to check", "Numbers end up in spreadsheets. It is difficult to see later where a quantity came from."],
          ["Revisions repeat the work", "When the consultant issues new drawings, much of the measuring is done again."],
          ["Know-how lives in people's heads", "How Static Group measures (heights, depths, allowances) is not written down anywhere."],
          ["Management sees little", "It is hard to know quickly which projects have problems or missing information."],
          ["The goal", "Let the computer do the repetitive measuring, show the proof, and let estimators decide."]], cols=3, h=Inches(2.0), size=14)
notes(s, "These are the everyday pains of preparing a BOQ manually. The system targets each one, without removing the estimator from the process.")

# 3 one-picture overview
s = base("How it works — the whole system in one picture", "Overview")
flow(s, [("Drawings", "CAD files (DXF) of the project"), ("Read", "Find layers, text, tables"), ("Recognise", "Footings, columns, beams, slabs, walls"),
         ("Measure", "Dimensions from drawings & schedules"), ("Calculate", "Concrete, blocks, plaster…"), ("Static Group BOQ", "Your own format and wording")], y=Inches(1.9))
flow(s, [("Estimator", "Checks every line"), ("Accept / Correct", "Or reject / keep open"), ("Rules", "Approve how Static Group measures"),
         ("Recalculate", "Only what is affected"), ("Audit", "Who decided what, when"), ("Export", "Excel BOQ & reports")], y=Inches(3.65), color=BLUE)
bullets(s, ["Top row: what the computer does automatically.", "Bottom row: what the estimator does — the computer never makes these decisions."], y=Inches(5.35), size=16)
notes(s, "The black boxes are start and end points. The top row is automatic. The bottom row is the human control loop that makes the result trustworthy.")

# 4 golden rules
s = base("Four promises the system keeps", "Principles")
cards(s, [["1. It never guesses", "If a dimension is missing, the line says 'Missing' or 'Needs Review'. It never invents a number."],
          ["2. Every number has proof", "Each quantity links back to the drawing, the element, its dimensions and the formula used."],
          ["3. The estimator decides", "The computer proposes. Only an estimator can accept, correct or approve a measuring rule."],
          ["4. Data stays in-house", "Runs on Static Group's own computer. Drawings and BOQs are not sent to outside services."]], cols=2, h=Inches(2.1), size=16)
notes(s, "These four rules were designed in from the start. They are the reason the output can be trusted and audited.")

# 5 reading drawings
s = base("Step 1 — Reading the drawings", "Drawings")
bullets(s, ["Opens every CAD (DXF) file of a project — about 80 files per villa, all disciplines.",
            "Sorts them: structural, architectural, mechanical, electrical.",
            "Finds what matters on each drawing: layers, text labels, tables (schedules) and level marks.",
            "Keeps a drawing register: drawing number, title, revision, source file, upload time and status.",
            "  Numbers and titles are read from the file name and clearly marked 'parsed' so they can be checked.",
            "  Revisions are 'not recorded' until someone types them in from the title block — never guessed.",
            "Files that cannot be read are shown as failed, not hidden."], w=Inches(6.2), size=16)
shot(s, "viewer.jpg", Inches(7.0), Inches(1.75), Inches(5.75))
notes(s, "The system reads the same DXF files the consultant provides. It does not need PDFs or scans. Everything it reads is listed so nothing is silently skipped.")

# 6 recognising & measuring
s = base("Step 2 — Recognising and measuring each element", "Measurement")
bullets(s, ["Finds each footing, raft, blinding, column, beam, slab, pedestal and RC wall.",
            "Reads its ID label (e.g. F5, C12, T.B1) and its floor.",
            "Takes sizes from the drawing outline and from the drawing's own schedules (tables).",
            "Cross-checks: does the drawn size match the schedule? If not, it flags it.",
            "Example — footing F5 on the foundation plan:",
            "  1.60 m × 1.20 m from the outline, 0.40 m depth from the footing schedule",
            "  1.60 × 1.20 × 0.40 = 0.768 m³ — with the exact drawing, layer and table row recorded"], w=Inches(6.4), size=16)
cards(s, [["Ramzi villa", "247 elements found and measured"], ["Al-Kawari villa", "377 elements found and measured"], ["Al-Sulaiti villa", "348 elements found and measured"]], cols=1, y=Inches(1.85), h=Inches(1.1), size=15, x0=Inches(7.4), total_w=Inches(5.3))
notes(s, "This is the core measuring engine. It is deterministic: the same drawings always give the same numbers, and every number can be traced.")

# 7 quantities
s = base("Step 3 — Calculating the BOQ quantities", "Quantities")
cards(s, [["Concrete", "Footings, ground beams, neck columns, slab on grade, columns & walls, beams, slabs"],
          ["Blinding", "Under footings (and ground beams, if Static Group measures it that way)"],
          ["Earthworks", "Excavation, removal, backfill — needs the natural ground level"],
          ["Waterproofing", "Below-ground surfaces of footings and beams"],
          ["Block work & plaster", "Wall areas by thickness, internal / external faces"],
          ["Lintels", "Over doors and windows"],
          ["Manholes", "Counted from drainage and electrical drawings"],
          ["Tanks", "From tank sizes the estimator enters"],
          ["Boundary wall", "From length and typical section the estimator enters"]], cols=3, h=Inches(1.35), size=13)
notes(s, "Quantities that depend on a Static Group habit (for example how column height is measured) are shown as options until an estimator approves the rule.")

# 8 Static Group BOQ
s = base("Step 4 — The BOQ in Static Group's own format", "BOQ")
bullets(s, ["Same sections, codes and wording as Static Group's real BOQs (learned from Al-Kawari and Al-Sulaiti).",
            "57 standard BOQ items are known to the system.",
            "Every line shows: code, description, unit, quantity, status, confidence, source and review state.",
            "Lines are never filled with a number that is not supported — blank means 'not final yet'.",
            "Exports to Excel in the same structure.",
            "If a client's BOQ has a line the system does not know, it is flagged as 'New BOQ item detected'."], w=Inches(6.0), size=16)
shot(s, "boq.jpg", Inches(6.85), Inches(1.75), Inches(5.9))
notes(s, "The output looks like a Static Group BOQ, not a generic list. The status column is the important part: it tells you how far you can trust each line.")

# 9 status labels
s = base("Five labels tell you how far to trust a line", "Status")
cards(s, [["Calculated", "The measuring rule is confirmed by an estimator and all the information is there."],
          ["Estimated", "Uses a rule learned from past projects (or a temporary assumption) that no estimator has confirmed yet."],
          ["Needs Review", "A person must decide: the rule is open (several options shown) or a measured value must be checked."],
          ["Missing", "Information is not in the drawings — e.g. ground level or tank sizes must be entered."],
          ["Lump Sum / Excluded", "No measurement needed (e.g. preliminaries, design, permits)."]], cols=5, h=Inches(3.0), size=14,
      colors=[GREEN, BLUE, AMBER, RED, NEUTRAL])
txt(s, Inches(0.6), Inches(5.2), Inches(12), Inches(1.0), "Nothing moves to 'Calculated' by itself. An unknown or assumed rule can never produce a final quantity.", 17, INK, True)
notes(s, "Today, on the three reference projects, no measured line is final yet: they are waiting for the estimator's rule decisions. That is by design.")

# 10 evidence
s = base("“Where did this number come from?”", "Evidence panel")
bullets(s, ["Click any BOQ line to open its evidence on the right.",
            "You see: the quantity, the reason for its status, the rule used, the formula and the source drawings.",
            "Below that: every element that was measured, with its sizes.",
            "Elements missing a dimension are listed first — they are why the line needs review.",
            "'Drawing' opens the drawing and highlights the exact element."], w=Inches(5.6), size=16)
shot(s, "evidence.jpg", Inches(6.45), Inches(1.75), Inches(6.3))
notes(s, "This is the feature estimators usually like most: nothing is a black box. Every number can be checked in seconds.")

# 11 drawing viewer
s = base("The drawing viewer", "Evidence")
bullets(s, ["Shows the real CAD drawing inside the application — zoom, pan, fit to screen.",
            "Measured elements are coloured by type; the selected one is outlined.",
            "Side panel: the element's sizes, quantity, formula and layer.",
            "Layers can be switched on and off.",
            "Opens automatically on the right element when you click 'Drawing' in the BOQ."], w=Inches(5.6), size=16)
shot(s, "viewer.jpg", Inches(6.45), Inches(1.75), Inches(6.3))
notes(s, "Example on screen: footing F5 on the Al-Kawari foundation plan, 1.6 × 1.2 × 0.4 m = 0.768 m³.")

# 12 learning from past projects
s = base("Learning from Static Group's past projects", "Benchmark")
bullets(s, ["Al-Kawari and Al-Sulaiti have drawings AND Static Group's own priced BOQ.",
            "The system measures those drawings itself and compares, line by line, with Static Group's figures.",
            "The goal is not to copy old numbers — it is to discover how Static Group measures.",
            "Big differences are shown, not hidden or 'tuned away'. Examples today:",
            "  Al-Kawari ground beams +1.2 % (good match) · Al-Sulaiti footings −4.4 % (good match)",
            "  Al-Kawari footings −9.2 % · Al-Sulaiti ground beams −90.5 % (tie-beam depths not readable — under investigation)",
            "Each line explains 'why is this different?' using data only."], w=Inches(6.4), size=15)
shot(s, "benchmark.jpg", Inches(7.15), Inches(1.75), Inches(5.6))
notes(s, "The benchmark is evidence for investigation. Where the drawings are clear and the method matches, the system lands close to Static Group's figures.")

# 13 the 22 rules
s = base("The 22 measuring rules", "Methodology — turning habits into written rules")
bullets(s, ["Some Static Group practices are not written down. Examples:",
            "  How is column height measured — floor to floor, or to the underside of the slab?",
            "  Is beam depth measured full, or below the slab?",
            "  How far does blinding project beyond the concrete?",
            "  How much working space is added around excavation?",
            "The system lists these as 22 rules. Each rule has a status:"], w=Inches(12), size=16)
cards(s, [["CONFIRMED · 0", "Approved by an estimator"], ["INFERRED · 3", "Seen in past projects, not yet confirmed"],
          ["ASSUMED · 1", "Temporary assumption (storey level)"], ["UNKNOWN · 18", "Not decided — options shown, no value chosen"]], cols=4, y=Inches(4.6), h=Inches(1.25), size=15,
      colors=[GREEN, BLUE, AMBER, RED])
notes(s, "This is the single most important input Static Group can give: an estimator spending about an hour deciding these 22 rules.")

# 14 rule approval screen
s = base("The Rule Approval screen", "Methodology")
bullets(s, ["One card per rule: the question, the evidence from past projects and why it matters.",
            "Each option shows the quantity it would give on every project, next to Static Group's real figure.",
            "Buttons: Approve rule · Keep unconfirmed (only adds a note) · Add note · Reject / Reopen · View evidence.",
            "Approving asks for confirmation, then recalculates only the projects that use the rule.",
            "If options give the same number on a project, the card explains why (e.g. no slabs detected on Al-Sulaiti)."], w=Inches(5.7), size=15)
shot(s, "rules.jpg", Inches(6.55), Inches(1.75), Inches(6.2))
notes(s, "The system never selects an option itself. It shows the consequences of each choice so the estimator can decide quickly and confidently.")

# 15 sessions
s = base("Validation sessions — who decided what, and when", "Traceability")
bullets(s, ["The estimator starts a 'validation session' with their name before reviewing.",
            "Every decision in that session is tagged to it: rules confirmed, rejected or kept open, and BOQ line decisions.",
            "Finishing the session freezes a summary: what was reviewed, what changed, which projects were recalculated.",
            "A session started by mistake can be cancelled — decisions already made are kept, nothing is erased.",
            "Only one session can be open at a time."], size=17)
notes(s, "Sessions make the validation traceable: later you can show exactly which estimator confirmed which rule, and what it changed.")

# 16 estimator line decisions
s = base("Checking every BOQ line", "Validation — accept, correct, reject or keep open")
cards(s, [["ACCEPT", "Approves the computer's quantity. Not possible when there is no computer quantity (e.g. rule still unknown)."],
          ["CORRECT", "The estimator types the right quantity and a short reason. The system stores both numbers and the difference."],
          ["REJECT", "The proposal is wrong. A reason is required. No quantity is kept."],
          ["KEEP UNCONFIRMED", "Not decided yet. Nothing becomes final."]], cols=4, y=Inches(1.7), h=Inches(1.55), size=13, colors=[GREEN, BLUE, RED, NEUTRAL])
shot(s, "validation.jpg", Inches(3.9), Inches(3.45), Inches(5.6), maxh=Inches(3.35))
notes(s, "Important: a correction never changes the measuring rule by itself. It becomes evidence that helps the estimator decide the rule later.")

# 17 baselines & revisions
s = base("Baselines and drawing revisions", "Revisions")
bullets(s, ["A baseline is a frozen set of quantities to compare against:",
            "  Static Group's existing BOQ, an imported client BOQ, or quantities the estimator approved.",
            "  Once saved it can never be changed — each baseline keeps its revision, source and date.",
            "When new drawings arrive, the system compares the baseline with the new result, line by line:",
            "  baseline quantity · new quantity · difference · estimator's quantity · status · rule · drawing revision",
            "  Clearly lists what was ADDED, OMITTED, CHANGED, or cannot be compared yet.",
            "If the computer's quantity changes after a decision, that decision is marked 'stale' and must be re-checked.",
            "Honest status: built and tested, but not yet validated on a real revised drawing set."], size=16)
notes(s, "Revision handling is where automation can save the most time later. It still needs one real revised drawing set to prove itself.")

# 18 project inputs
s = base("Information the drawings cannot give", "Inputs")
bullets(s, ["Some figures are not in the CAD files in a readable form. The estimator enters them once per project:",
            "  Natural ground level (needed for excavation and backfill)",
            "  Tank sizes (length, width, height, wall / base / cover thickness)",
            "  Boundary wall length and typical section",
            "Everything entered is labelled 'User input' on the quantities it affects and is recorded in the audit trail.",
            "If a value is not known, leave it empty — the lines stay 'Missing'. The system never fills the gap itself."], size=17)
notes(s, "Today these inputs are missing on all three reference projects, which is why those lines show as Missing.")

# 19 review queue
s = base("The review queue — the engineer's to-do list", "Review")
bullets(s, ["Collects everything that needs a person, in priority order:",
            "  1. Missing dimensions  2. Drawing conflicts  3. Missing project inputs",
            "  4. Large differences from past BOQs  5. Unconfirmed measuring rules  6. Other checks",
            "Each item shows the problem, the element, the reason, the evidence and the impact.",
            "Actions: Approve · Resolve · Reject · Reopen — and Correct where a value can be entered (e.g. slab thickness).",
            "Every action is recorded with name, date and comment."], size=17)
notes(s, "Examples of review items: a beam with no depth in the schedule, a label that matches two marks, a slab without a thickness.")

# 20 dashboard
s = base("Management dashboard", "Management — the whole picture in 30 seconds")
bullets(s, ["Real figures only: projects, BOQ lines, items requiring review, missing information, pending rules.",
            "Project overview with a progress bar for each project.",
            "'Attention required' list, most important first.",
            "'Automation readiness' table: why quantities are not final yet — no fake percentages.",
            "Time or cost savings are deliberately not shown until they are measured."], w=Inches(5.4), size=15)
shot(s, "dashboard.jpg", Inches(6.25), Inches(1.75), Inches(6.5))
notes(s, "Current figures: 3 projects, 237 BOQ lines, 147 items requiring review, 66 missing information, 22 pending rules.")

# 21 assistant
s = base("The project assistant (AI)", "Intelligence")
bullets(s, ["Ask questions in plain English, for example:",
            "  Which projects need attention?  ·  What information is missing?",
            "  Why is beam concrete different?  ·  Which BOQ items need review?  ·  What changed?",
            "Answers come only from the application's own data — with the sources listed underneath.",
            "It never measures drawings and never answers from general memory.",
            "If it cannot answer from the data, it says so instead of guessing.",
            "Works offline by default; no project data leaves the computer."], w=Inches(5.6), size=15)
shot(s, "assistant.jpg", Inches(6.45), Inches(1.75), Inches(6.3))
notes(s, "The AI helps people find and understand information. It does not calculate quantities and does not make decisions.")

# 22 daily report
s = base("Daily management report", "Intelligence")
bullets(s, ["One page, readable in under a minute:",
            "  critical items · items needing attention · pending decisions · quantity changes",
            "  recommended actions, large differences from past BOQs, missing information per project",
            "Can be saved as a snapshot each day, so new and resolved issues are shown the next day.",
            "Can be printed, or produced automatically every morning on a schedule."], w=Inches(5.6), size=16)
shot(s, "report.jpg", Inches(6.45), Inches(1.75), Inches(6.3))
notes(s, "No e-mail or Teams delivery is connected yet; that would need Static Group's approval of where project data may go.")

# 23 pilot measurement
s = base("Measuring the real value", "Pilot record")
bullets(s, ["For the first live project we record, side by side:",
            "  Manual: take-off hours, BOQ preparation hours, review hours, missing items found, corrections",
            "  System: processing time, calculation time, BOQ generation time, estimator review hours",
            "  Quantities: manual vs system for chosen BOQ items",
            "  Estimator decisions: accepted, corrected, rejected, still open — and the size of corrections",
            "A difference is only shown when both numbers exist.",
            "No accuracy %, savings or ROI are claimed until real estimator data proves them.",
            "All early figures are clearly labelled 'Preliminary'."], size=16)
notes(s, "This is how Static Group will get an honest business case: from its own estimators' numbers on a real tender.")

# 24 audit & data safety
s = base("Audit trail and data safety", "Control")
cards(s, [["Everything is recorded", "Rule decisions, line decisions, inputs, corrections, review actions, exports — with name, date, comment and before/after values."],
          ["Nothing is overwritten", "Baselines are frozen. Old processing runs are kept. Rejected or reopened rules keep their history."],
          ["Stays in-house", "Runs on Static Group's computer. Drawings and BOQs are not sent to outside services."],
          ["Test data separated", "Test projects are archived: kept on disk, but excluded from dashboards, reports and the assistant."]], cols=2, h=Inches(2.1), size=15)
notes(s, "For every number in a BOQ you can answer: which drawing, which rule, who approved it and when.")

# 25 where we are
s = base("Where we are today", "Status")
cards(s, [["3 reference projects", "Ramzi, Al-Kawari, Al-Sulaiti — about 80 drawings each, all processed"],
          ["237 BOQ lines", "79 · 82 · 76 lines. 24 lump sum / excluded. No measured line is final yet — by design."],
          ["22 measuring rules", "18 unknown · 3 inferred · 1 assumed · 0 confirmed — waiting for the estimator"],
          ["Best matches so far", "Al-Kawari ground beams +1.2 %, Al-Sulaiti footings −4.4 %"],
          ["Known gaps", "Annex buildings, parapets, door/window sizes, Al-Sulaiti slab boundaries and tie-beam depths"],
          ["Quality checks", "105 automatic tests passing; screens checked on desktop sizes; production data untouched by testing"]], cols=3, h=Inches(2.1), size=14)
notes(s, "The honest summary: the machinery is built and tested. What remains is Static Group's own knowledge and a real project.")

# 26 what we need
s = base("What we need from Static Group", "Next")
bullets(s, ["About one hour of an estimator's time to decide the 22 measuring rules.",
            "One real, unseen tender project: its drawing set (DXF) and, if available, the existing BOQ.",
            "Drawing revision letters from the title blocks.",
            "Ground level, tank sizes and boundary-wall section for that project.",
            "A manual take-off of the same project (with hours) to compare against.",
            "Later: a revised drawing set, to prove the revision comparison."], size=18)
notes(s, "None of these can be produced by the system — they are Static Group's knowledge and data.")

# 27 first live test
s = base("The first live test — step by step", "Next")
flow(s, [("1. New project", "Upload drawings & BOQ"), ("2. Process", "Read & measure"), ("3. Register", "Enter revisions & inputs"), ("4. Baseline", "Freeze the existing BOQ"),
         ("5. Session", "Estimator starts session")], y=Inches(1.85), h=Inches(1.15))
flow(s, [("6. Rules", "Decide the 22 rules"), ("7. Lines", "Accept / correct / reject / keep"), ("8. Pilot", "Enter manual hours & quantities"), ("9. Compare", "Baseline vs system"),
         ("10. Finish", "Close session, export BOQ & audit")], y=Inches(3.35), h=Inches(1.15), color=BLUE)
bullets(s, ["Capture: decision counts, size of corrections, missing information found, time spent — manual and automated."], y=Inches(4.95), size=16)
notes(s, "The whole test stays under the estimator's control. The output is a measured, honest comparison.")

# 28 status & close
page[0] += 1
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = INK
s.shapes.add_picture(LOGO_WHITE, Inches(0.8), Inches(0.7), height=Inches(0.7))
txt(s, Inches(0.8), Inches(1.9), Inches(11.5), Inches(0.8), "Status", 18, RGBColor(0x9A, 0x9F, 0xA6), True)
txt(s, Inches(0.8), Inches(2.4), Inches(12.0), Inches(0.9), "Ready for a controlled live test with a real project", 30, WHITE, True)
txt(s, Inches(0.8), Inches(3.4), Inches(11.8), Inches(2.4), [
    "Not yet ready for unsupervised production use, because:",
    "•  the 22 measuring rules still need an estimator's decision",
    "•  project inputs (ground level, tanks, boundary wall) and drawing revisions are still missing",
    "•  some extraction gaps remain, and revision comparison is not yet proven on real revised drawings",
    "•  no real manual-vs-system comparison has been measured yet"], 17, RGBColor(0xD5, 0xD8, 0xDC))
txt(s, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.6), "The computer proposes. The estimator decides. Every number shows its proof.", 18, WHITE, True)
notes(s, "Close by asking for the two things that unlock the next step: an estimator's hour for the rules, and a real tender project.")

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
