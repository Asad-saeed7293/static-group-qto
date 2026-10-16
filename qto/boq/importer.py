"""BOQ import (XLSX / CSV) with header detection and user-confirmable column mapping.

The importer never assumes column names: it proposes a mapping from synonyms,
the user confirms or changes it, and only then are rows normalised.
"""
from __future__ import annotations

import csv
import io
import os
import re

FIELDS = ["item_no", "description", "element", "floor", "quantity", "unit", "rate"]
REQUIRED = ["quantity"]
SYNONYMS = {
    "item_no": ["item no", "item", "item #", "ref", "ref.", "no", "no.", "s/n", "sn", "sl", "sl no", "boq ref", "code"],
    "description": ["description", "item description", "desc", "particulars", "work description", "scope"],
    "element": ["element", "element mark", "mark", "element id", "tag", "ref mark", "type", "member", "element ref"],
    "floor": ["floor", "level", "storey", "location", "zone"],
    "quantity": ["qty", "quantity", "boq qty", "boq quantity", "quant", "qty.", "amount qty"],
    "unit": ["unit", "uom", "units", "u.o.m", "unit of measure"],
    "rate": ["rate", "unit rate", "rate (qr)", "rate qr", "price", "unit price"],
}


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def read_table(path: str, sheet: str | None = None) -> dict:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheets = wb.sheetnames
        ws = wb[sheet] if sheet and sheet in sheets else wb[sheets[0]]
        rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return {"sheets": sheets, "sheet": ws.title, "rows": rows}
    if ext in (".csv", ".txt"):
        with open(path, "rb") as fh:
            raw = fh.read()
        text = raw.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect)]
        return {"sheets": ["CSV"], "sheet": "CSV", "rows": rows}
    raise ValueError(f"Unsupported BOQ file type '{ext}'. Use .xlsx or .csv")


def detect_header(rows: list) -> int:
    best, best_score = 0, -1
    for i, r in enumerate(rows[:25]):
        cells = [_norm(c) for c in r if c not in (None, "")]
        if len(cells) < 3:
            continue
        score = sum(1 for c in cells for syns in SYNONYMS.values() if c in syns)
        if score > best_score:
            best, best_score = i, score
    return best


def suggest_mapping(header: list) -> dict:
    mapping = {}
    used = set()
    for field in FIELDS:
        for i, h in enumerate(header):
            if i in used:
                continue
            if _norm(h) in SYNONYMS[field]:
                mapping[field] = i
                used.add(i)
                break
    for field in FIELDS:  # fuzzy: header contains a synonym
        if field in mapping:
            continue
        for i, h in enumerate(header):
            if i in used:
                continue
            if any(s in _norm(h) for s in SYNONYMS[field] if len(s) > 3):
                mapping[field] = i
                used.add(i)
                break
    return mapping


def preview(path: str, sheet: str | None = None) -> dict:
    t = read_table(path, sheet)
    rows = t["rows"]
    h = detect_header(rows)
    header = [str(c) if c is not None else "" for c in (rows[h] if rows else [])]
    title_rows = [" ".join(str(c) for c in r if c not in (None, "")) for r in rows[:h]]
    return {"sheets": t["sheets"], "sheet": t["sheet"], "header_row": h, "header": header,
            "suggested_mapping": suggest_mapping(header), "title_rows": title_rows,
            "sample": [[("" if c is None else str(c)) for c in r] for r in rows[h + 1:h + 9]],
            "row_count": max(0, len(rows) - h - 1)}


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return "INVALID:" + str(v)


def import_rows(path: str, mapping: dict, header_row: int, sheet: str | None = None) -> tuple[list, list]:
    missing = [f for f in REQUIRED if f not in mapping or mapping[f] in (None, "")]
    if missing:
        raise ValueError(f"Mapping must include: {', '.join(missing)}")
    rows = read_table(path, sheet)["rows"]
    out, warnings = [], []
    for ri, r in enumerate(rows[header_row + 1:], start=header_row + 2):
        def get(field):
            idx = mapping.get(field)
            if idx in (None, "") or int(idx) >= len(r):
                return None
            return r[int(idx)]
        if all(c in (None, "") for c in r):
            continue
        q = _num(get("quantity"))
        desc = get("description")
        el = get("element")
        if q is None and not desc and not el:
            continue  # section heading / blank line
        rec = {"row": ri, "item_no": str(get("item_no") or "").strip(), "description": str(desc or "").strip(),
               "element": str(el or "").strip(), "floor": str(get("floor") or "").strip(),
               "unit": str(get("unit") or "").strip(), "quantity": None, "rate": None, "issues": []}
        if isinstance(q, str):
            rec["issues"].append(f"Quantity '{q[8:]}' is not a number")
            warnings.append(f"Row {ri}: quantity '{q[8:]}' is not a number")
        else:
            rec["quantity"] = q
        if rec["quantity"] is None and not rec["issues"]:
            rec["issues"].append("BOQ quantity is blank")
        if "rate" in mapping:
            r = _num(get("rate"))
            if isinstance(r, str):
                rec["issues"].append(f"Rate '{r[8:]}' is not a number")
            elif r is not None and r > 0:
                rec["rate"] = r
        out.append(rec)
    return out, warnings
